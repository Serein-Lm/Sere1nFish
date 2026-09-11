"""Render Markdown structure into Word; used by both Word and PDF artifacts."""
from __future__ import annotations

from dataclasses import dataclass, replace
from typing import Any
from urllib.parse import urlsplit

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from markdown_it import MarkdownIt
from markdown_it.tree import SyntaxTreeNode


_CJK_FONT = "Noto Sans CJK SC"


def configure_document(doc: Any) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name in ("Normal", "Title", "Subtitle", "Quote", "List Paragraph", *[f"Heading {level}" for level in range(1, 10)]):
        style = doc.styles[name]
        style.font.name = _CJK_FONT
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), _CJK_FONT)
        if name.startswith("Heading") or name == "Title":
            style.font.color.rgb = RGBColor.from_string("203040")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2


@dataclass(frozen=True)
class _InlineStyle:
    bold: bool | None = None
    italic: bool | None = None
    strike: bool | None = None
    link: bool = False


def _add_run(paragraph: Any, text: str, style: _InlineStyle, *, parent: Any = None, code: bool = False) -> None:
    if not text:
        return
    run = paragraph.add_run(text)
    run.bold, run.italic, run.font.strike = style.bold, style.italic, style.strike
    if code:
        run.font.name = "DejaVu Sans Mono"
        run._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), _CJK_FONT)
        run.font.size = Pt(9)
    if style.link:
        run.font.color.rgb = RGBColor.from_string("1766A3")
        run.underline = True
    if parent is not None:
        parent.append(run._r)


def _hyperlink(paragraph: Any, destination: str) -> Any | None:
    from docx.opc.constants import RELATIONSHIP_TYPE

    try:
        url = urlsplit(destination)
    except ValueError:
        return None
    if url.scheme not in {"https", "http", "mailto"}:
        return None
    if url.scheme in {"https", "http"} and not url.hostname:
        return None
    element = OxmlElement("w:hyperlink")
    relation_id = paragraph.part.relate_to(destination, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    element.set(qn("r:id"), relation_id)
    paragraph._p.append(element)
    return element


def render_inline(paragraph: Any, nodes: list[SyntaxTreeNode], style: _InlineStyle | None = None, *, parent: Any = None) -> None:
    style = style or _InlineStyle()
    formatting = {"strong": "bold", "em": "italic", "s": "strike"}
    for node in nodes:
        if node.type in formatting:
            render_inline(paragraph, node.children, replace(style, **{formatting[node.type]: True}), parent=parent)
        elif node.type == "link":
            link = _hyperlink(paragraph, str(node.attrs.get("href") or ""))
            render_inline(paragraph, node.children, replace(style, link=link is not None), parent=link)
        elif node.type == "image":
            _add_run(paragraph, f"[图片：{node.content or '未提供说明'}]", style, parent=parent)
        elif node.children:
            render_inline(paragraph, node.children, style, parent=parent)
        elif node.type in {"softbreak", "hardbreak"}:
            _add_run(paragraph, "\n", style, parent=parent)
        else:
            _add_run(paragraph, node.content, style, parent=parent, code=node.type == "code_inline")


def _numbering(doc: Any, *, ordered: bool, start: int) -> int:
    root = doc.part.numbering_part.element
    ids = [int(value) for value in root.xpath("./w:abstractNum/@w:abstractNumId")]
    abstract_id = max(ids, default=-1) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multilevel = OxmlElement("w:multiLevelType")
    multilevel.set(qn("w:val"), "multilevel")
    abstract.append(multilevel)
    for level in range(9):
        definition = OxmlElement("w:lvl")
        definition.set(qn("w:ilvl"), str(level))
        values = {
            "start": str(start),
            "numFmt": "decimal" if ordered else "bullet",
            "lvlText": f"%{level + 1}." if ordered else ("•", "◦", "▪")[level % 3],
            "lvlJc": "left",
        }
        for key, value in values.items():
            element = OxmlElement(f"w:{key}")
            element.set(qn("w:val"), value)
            definition.append(element)
        properties = OxmlElement("w:pPr")
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), str((level + 1) * 720))
        indent.set(qn("w:hanging"), "360")
        properties.append(indent)
        definition.append(properties)
        abstract.append(definition)
    root.insert_element_before(abstract, "w:num")
    return root.add_num(abstract_id).numId


@dataclass(frozen=True)
class _BlockContext:
    list_depth: int = 0
    list_level: int = 0
    num_id: int | None = None
    quote_depth: int = 0


class DocxMarkdownRenderer:
    def __init__(self, doc: Any):
        self.doc = doc
        self.parser = MarkdownIt("commonmark", {"html": False}).enable(["table", "strikethrough"])
        self.handlers = {
            "paragraph": self._paragraph,
            "heading": self._heading,
            "table": self._table,
            "bullet_list": self._list,
            "ordered_list": self._list,
            "blockquote": self._quote,
            "fence": self._code,
            "code_block": self._code,
            "hr": self._rule,
        }

    def render(self, content: str) -> None:
        for node in SyntaxTreeNode(self.parser.parse(content)).children:
            self._block(node, _BlockContext())

    def _block(self, node: SyntaxTreeNode, context: _BlockContext) -> None:
        handler = self.handlers.get(node.type)
        if handler:
            handler(node, context)
        else:
            for child in node.children:
                self._block(child, context)

    def _new_paragraph(self, context: _BlockContext, *, style: str | None = None) -> Any:
        paragraph = self.doc.add_paragraph(style=style or ("Quote" if context.quote_depth else "Normal"))
        if context.num_id is not None:
            properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
            properties.get_or_add_numId().val = context.num_id
            properties.get_or_add_ilvl().val = min(context.list_level, 8)
        elif context.list_depth or context.quote_depth:
            paragraph.paragraph_format.left_indent = Cm(1.27 * context.list_depth + 0.6 * context.quote_depth)
        return paragraph

    def _paragraph(self, node: SyntaxTreeNode, context: _BlockContext) -> None:
        render_inline(self._new_paragraph(context), node.children)

    def _heading(self, node: SyntaxTreeNode, context: _BlockContext) -> None:
        paragraph = self._new_paragraph(context, style=f"Heading {int(node.tag[1:])}")
        render_inline(paragraph, node.children)

    def _list(self, node: SyntaxTreeNode, context: _BlockContext) -> None:
        number_id = _numbering(self.doc, ordered=node.type == "ordered_list", start=int(node.attrs.get("start", 1)))
        for item in node.children:
            first = True
            for child in item.children:
                child_context = replace(context, list_depth=context.list_depth + 1, list_level=context.list_depth, num_id=number_id if first else None)
                self._block(child, child_context)
                if child.type in {"paragraph", "heading"}:
                    first = False

    def _quote(self, node: SyntaxTreeNode, context: _BlockContext) -> None:
        for child in node.children:
            self._block(child, replace(context, quote_depth=context.quote_depth + 1))

    def _code(self, node: SyntaxTreeNode, context: _BlockContext) -> None:
        paragraph = self._new_paragraph(context)
        paragraph.paragraph_format.line_spacing = 1
        _add_run(paragraph, node.content.rstrip("\n"), _InlineStyle(), code=True)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F3F4F6")
        paragraph._p.get_or_add_pPr().append(shading)

    def _rule(self, node: SyntaxTreeNode, context: _BlockContext) -> None:
        paragraph = self._new_paragraph(context)
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in {"val": "single", "sz": "6", "color": "D9DEE5"}.items():
            bottom.set(qn(f"w:{key}"), value)
        borders.append(bottom)
        paragraph._p.get_or_add_pPr().append(borders)

    def _table(self, node: SyntaxTreeNode, context: _BlockContext) -> None:
        rows = [row for group in node.children for row in group.children]
        column_count = max((len(row.children) for row in rows), default=0)
        if not column_count:
            return
        table = self.doc.add_table(rows=len(rows), cols=column_count)
        table.style = "Table Grid"
        table.autofit = False
        section = self.doc.sections[-1]
        width = (section.page_width - section.left_margin - section.right_margin) // column_count
        for column in table.columns:
            column.width = width
        for index, row in enumerate(rows):
            self._table_row(table.rows[index], row, width=width)
        self.doc.add_paragraph()

    def _table_row(self, row: Any, node: SyntaxTreeNode, *, width: int) -> None:
        is_header = any(cell.type == "th" for cell in node.children)
        if is_header:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        alignments = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
        for index, source in enumerate(node.children):
            cell = row.cells[index]
            cell.width = width
            paragraph = cell.paragraphs[0]
            alignment = str(source.attrs.get("style", "")).partition(":")[2].strip()
            paragraph.alignment = alignments.get(alignment, WD_ALIGN_PARAGRAPH.LEFT)
            paragraph.paragraph_format.space_after = Pt(4)
            render_inline(paragraph, source.children, _InlineStyle(bold=is_header))
            if is_header:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EAF0F5")
                cell._tc.get_or_add_tcPr().append(shading)
