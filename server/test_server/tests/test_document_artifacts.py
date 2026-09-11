"""Document contracts: structured Word output and the same content in PDF."""
from io import BytesIO
import shutil

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
import pytest

from api.services.artifact_files import generate_artifact
from api.services.artifact_word import generate_docx


MARKDOWN = """## 中文章节

正文包含 **重点**、*强调*、~~删除~~ 和 `code()`。

| 项目 | 说明 | 数量 |
| :--- | :---: | ---: |
| 中文 | **完整** | 12 |
| 转义 | A\\|B | 3 |

3. 第三项
   - 嵌套子项
4. 第四项

> 引用内容

[参考来源](https://example.com/reference)

```python
items = [1, 2]
print(items)
```
"""


def test_word_renders_real_tables_inline_styles_and_chinese_fonts():
    result = generate_docx(title="排版验证", content=MARKDOWN)
    doc = Document(BytesIO(result["data"]))

    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3
    assert len(table.columns) == 3
    assert table.cell(1, 0).text == "中文"
    assert table.cell(1, 1).text == "完整"
    assert table.cell(2, 1).text == "A|B"
    assert table.cell(1, 1).paragraphs[0].runs[0].bold
    assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    assert int(table.cell(1, 2).paragraphs[0].alignment) == 2
    paragraph = next(p for p in doc.paragraphs if "正文包含" in p.text)
    assert next(r for r in paragraph.runs if r.text == "重点").bold
    assert next(r for r in paragraph.runs if r.text == "强调").italic
    assert next(r for r in paragraph.runs if r.text == "删除").font.strike
    assert doc.styles["Normal"].element.rPr.rFonts.get(qn("w:eastAsia"))
    assert any(p.style.name == "Heading 2" and p.text == "中文章节" for p in doc.paragraphs)
    assert not any("| 项目 |" in p.text for p in doc.paragraphs)


def test_word_preserves_list_start_nested_lists_code_and_links():
    doc = Document(BytesIO(generate_docx(title="结构验证", content=MARKDOWN)["data"]))
    numbered = next(p for p in doc.paragraphs if p.text == "第三项")
    nested = next(p for p in doc.paragraphs if p.text == "嵌套子项")
    assert numbered._p.pPr.numPr is not None
    assert nested._p.pPr.numPr.ilvl.val == 1
    assert any(value.get(qn("w:val")) == "3" for value in doc.part.numbering_part.element.iter(qn("w:start")))
    assert any("items = [1, 2]\nprint(items)" in p.text for p in doc.paragraphs)
    assert any(rel.target_ref == "https://example.com/reference" for rel in doc.part.rels.values())


def test_word_sections_use_the_same_markdown_renderer_without_fetching_images():
    result = generate_docx(
        title="章节验证",
        content="开场说明",
        sections=[{"heading": "数据", "body": "|列|值|\n|---|---|\n|甲|乙|\n\n![图片说明](http://127.0.0.1/private.png)"}],
    )
    doc = Document(BytesIO(result["data"]))
    assert len(doc.tables) == 1
    assert doc.tables[0].cell(1, 1).text == "乙"
    assert any("图片说明" in p.text for p in doc.paragraphs)
    assert not doc.inline_shapes


def test_pdf_registry_reuses_word_layout(monkeypatch):
    from core import office

    converted = []

    def convert(data, *, suffix, **kwargs):
        document = Document(BytesIO(data))
        converted.append(document.tables[0].cell(1, 0).text)
        assert suffix == ".docx"
        return b"%PDF-test"

    monkeypatch.setattr(office, "convert_office_to_pdf", convert)
    result = generate_artifact(title="中文附件", content=MARKDOWN, output_format=".PDF")
    assert converted == ["中文"]
    assert result["kind"] == "pdf"
    assert result["content_type"] == "application/pdf"
    assert result["filename"] == "中文附件.pdf"
    assert result["data"].startswith(b"%PDF")


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice is required")
def test_pdf_opens_and_contains_the_word_content():
    result = generate_artifact(title="Document check", content=MARKDOWN, output_format="pdf")
    reader = PdfReader(BytesIO(result["data"]))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert len(reader.pages) >= 1
    assert "Document check" in text
    assert "12" in text
    assert "A|B" in text
    assert result["size"] == len(result["data"])


def test_pdf_tool_returns_a_downloadable_artifact(monkeypatch):
    from Sere1nGraph.graph.tools import word_tools
    from core import office

    saved = []
    monkeypatch.setattr(office, "convert_office_to_pdf", lambda *args, **kwargs: b"%PDF-test")

    def persist(result, *, meta):
        saved.append((result, meta))
        return {**result, "download_url": f"/api/v1/artifacts/{result['artifact_id']}/download"}

    monkeypatch.setattr(word_tools, "_persist_artifact", persist)
    response = word_tools.generate_document_artifact.invoke({"title": "附件", "content": "内容", "output_format": "pdf"})
    assert "PDF" in response
    assert "[[artifact:art_" in response
    assert "/download" in response
    assert saved[0][1]["output_format"] == "pdf"


def test_pdf_failure_does_not_register_an_artifact(monkeypatch):
    from Sere1nGraph.graph.tools import word_tools
    from core import office

    def unavailable(*args, **kwargs):
        raise RuntimeError("运行环境缺少 LibreOffice")

    monkeypatch.setattr(office, "convert_office_to_pdf", unavailable)
    monkeypatch.setattr(word_tools, "_persist_artifact", lambda *args, **kwargs: pytest.fail("failed conversion must not persist"))
    response = word_tools.generate_document_artifact.invoke({"title": "附件", "content": "内容", "output_format": "pdf"})
    assert "失败" in response and "LibreOffice" in response
